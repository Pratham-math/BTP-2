"""Bounded parsers. No extraction is silently counted as author confirmation."""
import re
import zipfile
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from server.lessons.contracts import EvidenceBlock
from server.providers.base import GenerationError


def text_blocks(text, *, method="typed", location="Notes"):
    if not text.strip():
        raise GenerationError("EMPTY_SOURCE", "The source has no readable text.")
    # Split at paragraph boundaries while preserving indentation inside code.
    paragraphs = re.split(r"\n\s*\n", text.replace("\r\n", "\n"))
    chunks, current = [], ""
    for paragraph in paragraphs:
        if len(paragraph) > 20000:
            raise GenerationError("SOURCE_BLOCK_TOO_LARGE", "A source paragraph exceeds 20000 characters. Split it into sections before confirming.")
        if len(current) + len(paragraph) > 14000:
            chunks.append(current)
            current = ""
        current += ("\n\n" if current else "") + paragraph
    if current:
        chunks.append(current)
    return [EvidenceBlock(id=f"text-{i}", text=chunk, location=f"{location}, section {i}", method=method).model_dump() for i, chunk in enumerate(chunks, 1)]


def youtube_id(url):
    try:
        parsed = urlsplit(url.strip())
        valid = parsed.scheme == "https" and not parsed.username and not parsed.password and parsed.port in {None, 443}
    except ValueError:
        valid = False
    if not valid:
        raise GenerationError("UNSUPPORTED_URL", "Use an HTTPS public YouTube watch URL.")
    if parsed.hostname == "youtu.be":
        video_id = parsed.path.strip("/")
    elif parsed.hostname in {"www.youtube.com", "youtube.com", "m.youtube.com"}:
        if parsed.path == "/watch":
            video_id = parse_qs(parsed.query).get("v", [""])[0]
        elif re.fullmatch(r"/(?:shorts|live|embed)/[a-zA-Z0-9_-]{11}/?", parsed.path):
            video_id = parsed.path.strip("/").split("/")[1]
        else:
            video_id = ""
    else:
        raise GenerationError("UNSUPPORTED_URL", "Only public YouTube watch URLs are accepted.")
    if not re.fullmatch(r"[a-zA-Z0-9_-]{11}", video_id):
        raise GenerationError("UNSUPPORTED_URL", "The YouTube video ID is invalid.")
    return video_id


def validate_document(path, kind):
    with open(path, "rb") as stream:
        header = stream.read(12)
    if kind == "pdf" and not header.startswith(b"%PDF-"):
        raise GenerationError("INVALID_SOURCE", "The file is not a PDF.")
    if kind == "pptx":
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if (len(entries) > 10000 or sum(x.file_size for x in entries) > 100 * 1024 * 1024
                        or any(x.flag_bits & 1 or x.filename.startswith("/") or ".." in Path(x.filename).parts for x in entries)
                        or "ppt/presentation.xml" not in archive.namelist()
                        or any("vbaProject" in x.filename or "embeddings/" in x.filename for x in entries)):
                    raise ValueError("Unsafe archive")
        except (zipfile.BadZipFile, ValueError):
            raise GenerationError("INVALID_SOURCE", "Use a valid, unencrypted PPTX without macros or embedded programs.") from None


def extract_document(path, kind):
    validate_document(path, kind)
    blocks, warnings = [], []
    if kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise GenerationError("ENCRYPTED_SOURCE", "Upload an unencrypted PDF.")
        if len(reader.pages) > 200:
            raise GenerationError("SOURCE_BATCH_REQUIRED", "This PDF exceeds 200 pages. Split it into explicit batches; no pages have been dropped.")
        for i, page in enumerate(reader.pages, 1):
            value = page.extract_text(extraction_mode="layout") or ""
            uncertainty = []
            if not value.strip():
                value = "[No readable digital text; OCR and confirmation required.]"
                uncertainty = ["ocr_required"]
                warnings.append({"location": f"Page {i}", "code": "ocr_required", "critical": True})
            if len(value) > 20000:
                raise GenerationError("SOURCE_BATCH_REQUIRED", f"Page {i} needs splitting into smaller evidence blocks.")
            blocks.append(EvidenceBlock(id=f"page-{i}", text=value, location=f"Page {i}", method="pypdf-layout", uncertainty=uncertainty).model_dump())
        return blocks, warnings
    if kind == "pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        if len(presentation.slides) > 200:
            raise GenerationError("SOURCE_BATCH_REQUIRED", "This presentation exceeds 200 slides. Split it into explicit batches.")
        for i, slide in enumerate(presentation.slides, 1):
            lines, uncertain = [], []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines.append(shape.text)
                if shape.has_table:
                    lines.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                # Relationships and figures need a rendered preview/interpretation.
                if shape.shape_type in {6, 13} or getattr(shape, "has_chart", False):
                    uncertain.append("figure_confirmation_required")
            if slide.has_notes_slide:
                lines.append(slide.notes_slide.notes_text_frame.text)
            value = "\n\n".join(lines).strip() or "[Slide contains no readable text.]"
            if len(value) > 20000:
                raise GenerationError("SOURCE_BATCH_REQUIRED", f"Slide {i} needs splitting before import.")
            if not lines:
                uncertain.append("empty_slide")
            blocks.append(EvidenceBlock(id=f"slide-{i}", text=value, location=f"Slide {i}", method="python-pptx", uncertainty=uncertain).model_dump())
            warnings.extend({"location": f"Slide {i}", "code": flag, "critical": True} for flag in set(uncertain))
        return blocks, warnings
    raise GenerationError("UNSUPPORTED_SOURCE", "This document format needs a configured conversion or vision adapter.")
